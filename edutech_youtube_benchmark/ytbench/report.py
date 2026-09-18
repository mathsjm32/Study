"""5단계: 엑셀 리포트 생성.

기획자가 엑셀을 열어 바로 쓸 수 있는 형태를 목표로 한다.
  · 00_요약      : 한 장으로 끝나는 경영 보고용 요약
  · 01_채널랭킹   : 193개 정렬·색상 스케일 (의사결정의 기본 화면)
  · 02_성공패턴   : 업계 전체에서 검증된 제목/포맷 리프트
  · 03_LLM인사이트: 채널별 배울 점 / 부족한 점 / 실행안
  · 04~09        : 원자료, 상위영상, 타이밍 히트맵, 세그먼트, 검수, 지표정의
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .score import MIN_PATTERN_SAMPLE

log = logging.getLogger(__name__)

# 색상 — 채도를 낮춰 장시간 열람에도 눈이 피로하지 않게
NAVY = "1F3864"
ACCENT = "2E5C8A"
LIGHT = "EEF3F9"
WARN = "FFF2CC"
BAD = "FCE4E4"

HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(color=NAVY, bold=True, size=14)
SUB_FONT = Font(color=ACCENT, bold=True, size=11)
THIN = Side(style="thin", color="D0D7E2")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _write_frame(
    ws,
    df: pd.DataFrame,
    start_row: int = 1,
    freeze: str | None = None,
    widths: dict[str, int] | None = None,
    wrap_cols: set[str] | None = None,
) -> int:
    """DataFrame을 시트에 쓰고 헤더 서식을 적용한다. Returns: 마지막 행 번호."""
    if df.empty:
        ws.cell(row=start_row, column=1, value="(데이터 없음)")
        return start_row

    for c, col in enumerate(df.columns, 1):
        cell = ws.cell(row=start_row, column=c, value=str(col))
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    wrap_cols = wrap_cols or set()
    for r, (_, row) in enumerate(df.iterrows(), start_row + 1):
        for c, col in enumerate(df.columns, 1):
            value = row[col]
            if isinstance(value, float) and pd.isna(value):
                value = None
            elif hasattr(value, "item"):          # numpy 스칼라 -> 파이썬 기본형
                try:
                    value = value.item()
                except (ValueError, AttributeError):
                    value = str(value)
            elif isinstance(value, pd.Timestamp):
                value = value.tz_localize(None) if value.tzinfo else value
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = BORDER
            if col in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.row_dimensions[start_row].height = 30
    for c, col in enumerate(df.columns, 1):
        letter = get_column_letter(c)
        if widths and col in widths:
            ws.column_dimensions[letter].width = widths[col]
        else:
            # pandas 3.0의 astype(str)은 결측을 float nan으로 남기므로 map(str)을 쓴다
            sample = df[col].head(200).map(str)
            longest = max([len(str(col))] + [len(s) for s in sample])
            ws.column_dimensions[letter].width = min(max(10, longest + 2), 45)

    if freeze:
        ws.freeze_panes = freeze
    ws.auto_filter.ref = (
        f"A{start_row}:{get_column_letter(len(df.columns))}{start_row + len(df)}"
    )
    return start_row + len(df)


def _color_scale(ws, col_letter: str, first: int, last: int, reverse: bool = False) -> None:
    lo, hi = ("F8696B", "63BE7B") if not reverse else ("63BE7B", "F8696B")
    ws.conditional_formatting.add(
        f"{col_letter}{first}:{col_letter}{last}",
        ColorScaleRule(start_type="min", start_color=lo,
                       mid_type="percentile", mid_value=50, mid_color="FFEB84",
                       end_type="max", end_color=hi),
    )


# ── 시트별 작성 ───────────────────────────────────────────────────────
def _sheet_summary(wb, scores, videos, industry, lift, insights_df) -> None:
    ws = wb.create_sheet("00_요약")
    ws["A1"] = "에듀테크 경쟁사 유튜브 벤치마킹 리포트"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"분석일 {pd.Timestamp.now(tz=config.KST):%Y-%m-%d} · "
                f"채널 {len(scores)}개 · 영상 {len(videos):,}개")
    ws["A2"].font = Font(size=10, color="555555")

    row = 4
    ws.cell(row=row, column=1, value="1. 업계 기준선").font = SUB_FONT
    row = _write_frame(ws, industry, start_row=row + 1) + 2

    ws.cell(row=row, column=1, value="2. 상위 10개 채널").font = SUB_FONT
    top_cols = ["업계순위", "company_name", "등급", "종합점수", "subscriber_count",
                "median_views", "engagement_rate_median", "uploads_last_90d", "진단"]
    top = scores[[c for c in top_cols if c in scores.columns]].head(10).rename(columns={
        "company_name": "회사명", "subscriber_count": "구독자",
        "median_views": "중위조회수", "engagement_rate_median": "참여율",
        "uploads_last_90d": "90일업로드",
    })
    row = _write_frame(ws, top, start_row=row + 1, widths={"진단": 28}) + 2

    ws.cell(row=row, column=1, value="3. 검증된 성공 패턴 TOP 10 (리프트 순)").font = SUB_FONT
    if not lift.empty:
        row = _write_frame(ws, lift.head(10), start_row=row + 1) + 2
    else:
        ws.cell(row=row + 1, column=1, value="(표본 부족으로 미산출)")
        row += 3

    ws.cell(row=row, column=1, value="4. 읽는 방법").font = SUB_FONT
    guide = [
        "종합점수는 절대 성과가 아니라 '수집된 경쟁사 집단 내 상대 위치(백분위 가중합)'입니다.",
        "규모가 큰 채널이 자동으로 높은 점수를 받지 않도록, 효율성 축에 구독자당 지표를 25% 가중했습니다.",
        "리프트는 채널 내부 상대 조회수로 계산해, 대형 채널의 패턴이 과대평가되는 편향을 제거했습니다.",
        "'배울 점/부족한 점'은 03_LLM인사이트 시트에 채널별로 근거와 함께 정리돼 있습니다.",
        "매칭 신뢰도가 낮은 채널은 08_매칭검수 시트에서 먼저 확인하세요. 오매칭은 모든 수치를 왜곡합니다.",
    ]
    for i, line in enumerate(guide):
        c = ws.cell(row=row + 1 + i, column=1, value=f"· {line}")
        c.alignment = Alignment(wrap_text=False)

    ws.column_dimensions["A"].width = 34
    for letter in "BCDEFGHI":
        ws.column_dimensions[letter].width = 16


def _sheet_ranking(wb, scores) -> None:
    ws = wb.create_sheet("01_채널랭킹")
    cols = ["업계순위", "등급", "company_name", "channel_title", "category", "channel_url",
            "종합점수"] + [f"점수_{a}" for a in config.SCORE_WEIGHTS] + [
            "subscriber_count", "median_views", "views_per_subscriber",
            "engagement_rate_median", "uploads_last_90d", "upload_interval_median_days",
            "days_since_last_upload", "숏폼_비율", "momentum_90d", "videos_sampled",
            "match_confidence", "진단"]
    df = scores[[c for c in cols if c in scores.columns]].rename(columns={
        "company_name": "회사명", "channel_title": "채널명", "category": "세그먼트",
        "channel_url": "채널링크", "subscriber_count": "구독자수",
        "median_views": "중위조회수", "views_per_subscriber": "구독자당조회수",
        "engagement_rate_median": "참여율중위", "uploads_last_90d": "90일업로드수",
        "upload_interval_median_days": "업로드간격일", "days_since_last_upload": "최종업로드경과일",
        "momentum_90d": "모멘텀90일", "videos_sampled": "분석영상수",
        "match_confidence": "매칭신뢰도",
    })
    last = _write_frame(ws, df, freeze="D2", widths={"진단": 26, "채널링크": 20})

    headers = {str(c): get_column_letter(i) for i, c in enumerate(df.columns, 1)}
    for name in ["종합점수"] + [f"점수_{a}" for a in config.SCORE_WEIGHTS]:
        if name in headers:
            _color_scale(ws, headers[name], 2, last)
    if "구독자수" in headers:
        ws.conditional_formatting.add(
            f"{headers['구독자수']}2:{headers['구독자수']}{last}",
            DataBarRule(start_type="min", end_type="max", color=ACCENT),
        )
    # 최종업로드 경과일은 낮을수록 좋으므로 색 스케일을 뒤집는다
    if "최종업로드경과일" in headers:
        _color_scale(ws, headers["최종업로드경과일"], 2, last, reverse=True)


def _sheet_insights(wb, insights_df, scores) -> None:
    ws = wb.create_sheet("03_LLM인사이트")
    if insights_df.empty:
        ws["A1"] = "LLM 인사이트가 생성되지 않았습니다."
        ws["A2"] = "ANTHROPIC_API_KEY 를 설정한 뒤 `python -m ytbench llm` 을 실행하세요."
        ws.column_dimensions["A"].width = 70
        return

    rank = scores[["company_name", "업계순위", "등급", "종합점수"]]
    df = rank.merge(insights_df, on="company_name", how="right").sort_values(
        "업계순위", na_position="last"
    )
    wrap = {c for c in df.columns if any(
        k in c for k in ("배울점", "부족한점", "실행안", "포지셔닝", "성공공식", "근거", "기회", "이유", "오류")
    )}
    widths = {c: (46 if c in wrap else 12) for c in df.columns}
    widths["company_name"] = 18
    _write_frame(ws, df, freeze="B2", widths=widths, wrap_cols=wrap)


def _sheet_timing(wb, videos) -> None:
    """요일 × 시간대 히트맵 — 업로드 타이밍과 성과의 관계."""
    ws = wb.create_sheet("06_업로드타이밍")
    if videos.empty:
        ws["A1"] = "(데이터 없음)"
        return

    ws["A1"] = "업로드 타이밍 분석 (KST 기준)"
    ws["A1"].font = TITLE_FONT

    from .features import WEEKDAY_KR
    row = 3
    for label, values, fmt in (
        ("업로드 건수 (요일 × 시간대)", "count", "0"),
        ("채널내 상대조회수 중위 (요일 × 시간대)", "median", "0.00"),
    ):
        ws.cell(row=row, column=1, value=label).font = SUB_FONT
        pivot = videos.pivot_table(
            index="upload_weekday",
            columns="upload_hour",
            values="views_vs_channel_median",
            aggfunc="size" if values == "count" else "median",
        ).reindex(WEEKDAY_KR)
        pivot.index.name = "요일"
        flat = pivot.reset_index()
        flat.columns = ["요일"] + [f"{int(h)}시" for h in pivot.columns]
        last = _write_frame(ws, flat, start_row=row + 1)
        first_col, last_col = get_column_letter(2), get_column_letter(len(flat.columns))
        ws.conditional_formatting.add(
            f"{first_col}{row + 2}:{last_col}{last}",
            ColorScaleRule(start_type="min", start_color="FFFFFF",
                           end_type="max", end_color="2E75B6"),
        )
        for r in range(row + 2, last + 1):
            for c in range(2, len(flat.columns) + 1):
                ws.cell(row=r, column=c).number_format = fmt
        row = last + 3


def _sheet_segments(wb, scores) -> None:
    ws = wb.create_sheet("07_세그먼트비교")
    if scores.empty or "category" not in scores.columns:
        ws["A1"] = "(세그먼트 정보 없음 — 입력 CSV의 category 컬럼을 채우면 활성화됩니다)"
        ws.column_dimensions["A"].width = 70
        return

    df = scores.copy()
    df["category"] = df["category"].replace("", "미분류").fillna("미분류")
    agg = df.groupby("category", observed=True).agg(
        채널수=("company_name", "count"),
        종합점수_평균=("종합점수", "mean"),
        구독자_중위=("subscriber_count", "median"),
        중위조회수_중위=("median_views", "median"),
        참여율_중위=("engagement_rate_median", "median"),
        업로드90일_중위=("uploads_last_90d", "median"),
        숏폼비율_평균=("숏폼_비율", "mean"),
    ).round(4).reset_index().rename(columns={"category": "세그먼트"})
    agg = agg.sort_values("종합점수_평균", ascending=False)
    last = _write_frame(ws, agg, freeze="B2")
    _color_scale(ws, get_column_letter(3), 2, last)


def _sheet_review(wb, resolved: pd.DataFrame | None) -> None:
    ws = wb.create_sheet("08_매칭검수")
    ws["A1"] = "채널 매칭 검수 대상"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = ("아래 항목은 자동 매칭 신뢰도가 낮습니다. 채널을 직접 확인해 "
                "companies.csv 의 channel_url 컬럼을 채우고 재실행하세요.")
    ws["A2"].font = Font(size=10, color="C00000")

    if resolved is None or resolved.empty:
        ws["A4"] = "(매칭 결과 파일이 없습니다)"
        ws.column_dimensions["A"].width = 80
        return

    need = resolved[
        (resolved["needs_review"].astype(str).str.lower().isin(["true", "1", "yes"]))
        | (resolved["match_method"] == "failed")
    ]
    body = need if not need.empty else resolved.head(0)
    last = _write_frame(ws, body, start_row=4,
                        widths={"candidates": 50, "reason": 40, "company_name": 20})
    if not need.empty:
        for r in range(5, last + 1):
            for c in range(1, len(body.columns) + 1):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor=WARN)
    else:
        ws["A4"] = "검수 대상 없음 — 전체 채널이 신뢰도 기준을 통과했습니다."
    ws.column_dimensions["A"].width = 24


DEFINITIONS = [
    ("종합점수", "6축 백분위 가중합 (0~100)", "수집된 경쟁사 집단 내 상대 위치. 절대 성과가 아님"),
    ("점수_도달력", "구독자수·총조회수 백분위 평균", "가중치 15%. 순수 규모"),
    ("점수_효율성", "구독자당 조회수·중위조회수·일일조회속도", "가중치 25%. 규모 보정 후 실제 성과"),
    ("점수_참여도", "(좋아요+댓글)/조회수 중위 + 히트영상 비율", "가중치 20%. 반응의 깊이와 폭"),
    ("점수_일관성", "90일 업로드 수 + 업로드 간격 CV(역방향)", "가중치 15%. 예측 가능한 발행 리듬"),
    ("점수_콘텐츠설계", "리프트가 검증된 제목·후킹 패턴 활용도", "가중치 15%. 데이터로 검증된 패턴 기준"),
    ("점수_최신성", "최종 업로드 경과일(역방향) + 30일 업로드 수", "가중치 10%. 현재 운영 여부"),
    ("중위조회수", "표본 영상 조회수의 중위값", "평균은 히트 1편에 왜곡되므로 중위값을 주 지표로 사용"),
    ("구독자당조회수", "채널 총조회수 / 구독자수", "구독자를 얼마나 반복 시청으로 전환하는가"),
    ("참여율중위", "(좋아요+댓글)/조회수 의 중위값", "좋아요·댓글이 비공개인 영상은 계산에서 제외"),
    ("업로드간격일", "연속 업로드 간 일수의 중위값", "작을수록 발행이 잦음"),
    ("규칙성 CV", "업로드 간격의 표준편차/평균", "낮을수록 규칙적. 몰아올리기 채널은 높게 나옴"),
    ("모멘텀90일", "최근 90일 중위조회수 / 직전 90일 중위조회수", "1.0 초과 성장, 미만 둔화. 표본 부족 시 공란"),
    ("히트영상 비율", "채널 중위 조회수의 3배 이상인 영상 비율", "터지는 콘텐츠를 만들 확률"),
    ("views_vs_channel_median", "영상 조회수 / 소속 채널의 중위 조회수", "채널 규모를 제거한 상대 성과. 패턴 리프트의 기준값"),
    ("리프트", "패턴 적용 영상의 상대조회수 중위 / 미적용 중위", "1.0 초과면 해당 패턴이 성과에 기여. 최소 30개 영상 필요"),
    ("숏폼/미들폼/롱폼", "60초 이하 / 60초~10분 / 10분 초과", "포맷별 전략이 달라 반드시 분리 집계"),
    ("매칭신뢰도", "회사명↔채널명 유사도 0.8 + 교육관련성 0.2", "0.70 이상 자동확정, 0.45~0.70 검수 권장"),
]


def _sheet_definitions(wb) -> None:
    ws = wb.create_sheet("09_지표정의")
    ws["A1"] = "지표 정의서"
    ws["A1"].font = TITLE_FONT
    df = pd.DataFrame(DEFINITIONS, columns=["지표", "계산 방식", "해석 / 주의점"])
    _write_frame(ws, df, start_row=3, freeze="A4",
                 widths={"지표": 24, "계산 방식": 46, "해석 / 주의점": 60},
                 wrap_cols={"계산 방식", "해석 / 주의점"})


# ── 엔트리포인트 ──────────────────────────────────────────────────────
def build_report(
    scores: pd.DataFrame,
    channel_metrics: pd.DataFrame,
    video_features: pd.DataFrame,
    industry: pd.DataFrame,
    lift_table: pd.DataFrame,
    insights_df: pd.DataFrame | None = None,
    resolved: pd.DataFrame | None = None,
    path: Path = config.REPORT_XLSX,
) -> Path:
    insights_df = insights_df if insights_df is not None else pd.DataFrame()
    wb = Workbook()
    wb.remove(wb.active)

    _sheet_summary(wb, scores, video_features, industry, lift_table, insights_df)
    _sheet_ranking(wb, scores)

    ws = wb.create_sheet("02_성공패턴")
    ws["A1"] = "제목·포맷 패턴별 성과 리프트"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"리프트 1.0 초과 = 해당 패턴이 성과에 기여. "
                f"패턴별 최소 {MIN_PATTERN_SAMPLE}개 영상 표본 기준.")
    ws["A2"].font = Font(size=10, color="555555")
    if not lift_table.empty:
        last = _write_frame(ws, lift_table, start_row=4, freeze="A5")
        headers = {str(c): get_column_letter(i) for i, c in enumerate(lift_table.columns, 1)}
        _color_scale(ws, headers["리프트"], 5, last)
    else:
        ws["A4"] = "(표본 부족으로 미산출)"

    _sheet_insights(wb, insights_df, scores)

    ws = wb.create_sheet("04_채널지표전체")
    _write_frame(ws, channel_metrics, freeze="C2")

    ws = wb.create_sheet("05_상위영상")
    if not video_features.empty:
        cols = ["company_name", "title", "format", "duration_sec", "view_count",
                "views_vs_channel_median", "engagement_rate", "upload_weekday",
                "upload_hour", "published_kst", "thumbnail_url"]
        top = (video_features[[c for c in cols if c in video_features.columns]]
               .nlargest(300, "views_vs_channel_median")
               .rename(columns={"company_name": "회사명", "title": "제목",
                                "format": "포맷", "duration_sec": "길이초",
                                "view_count": "조회수",
                                "views_vs_channel_median": "채널중위대비",
                                "engagement_rate": "참여율",
                                "upload_weekday": "요일", "upload_hour": "시",
                                "published_kst": "업로드시각",
                                "thumbnail_url": "썸네일링크"}))
        _write_frame(ws, top, freeze="C2", widths={"제목": 55, "썸네일링크": 20})
    else:
        ws["A1"] = "(데이터 없음)"

    _sheet_timing(wb, video_features)
    _sheet_segments(wb, scores)
    _sheet_review(wb, resolved)
    _sheet_definitions(wb)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    log.info("리포트 저장: %s (시트 %d개)", path, len(wb.sheetnames))
    return path
